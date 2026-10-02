"""Инструменты для ИИ: создание файлов/архивов, поиск в интернете, чтение сайтов."""
import asyncio, csv, html as htmlmod, io, ipaddress, json, os, re, socket, time, uuid, zipfile
from datetime import datetime, timezone, timedelta
from urllib.parse import urlparse, parse_qs, quote

FILES_DIR = os.getenv("FILES_DIR", "/tmp/ai-studio-files")
FILE_TTL = 24 * 3600
MAX_TOTAL = 15 * 1024 * 1024
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36"

TOOLS = [
    {"type": "function", "function": {
        "name": "create_files",
        "description": ("Создаёт файлы и сразу отправляет их пользователю (в чат и в приложение). Используй ВСЕГДА, когда "
                        "просят создать/сохранить/скинуть файл, документ, код, таблицу, архив. Поддерживаются любые "
                        "текстовые форматы (.txt .md .py .js .html .csv .json и т.д.), а также .docx (содержимое в Markdown: "
                        "# заголовки, - списки, **жирный**) и .xlsx (содержимое в CSV; несколько листов — разделяй строкой "
                        "'### Sheet: Название'). Чтобы упаковать файлы в zip-архив — укажи archive_name. Пустые файлы — content: \"\"."),
        "parameters": {"type": "object", "properties": {
            "files": {"type": "array", "description": "Список файлов", "items": {"type": "object", "properties": {
                "path": {"type": "string", "description": "Имя файла, можно с папками внутри архива: src/main.py"},
                "content": {"type": "string", "description": "Содержимое файла"}}, "required": ["path", "content"]}},
            "archive_name": {"type": "string", "description": "Если задано — все файлы упаковываются в этот .zip архив"}},
            "required": ["files"]}}},
    {"type": "function", "function": {
        "name": "web_search",
        "description": "Поиск в интернете. Используй для свежих новостей, фактов, цен, курсов, событий после даты обучения и всего, в чём не уверен.",
        "parameters": {"type": "object", "properties": {"query": {"type": "string"}}, "required": ["query"]}}},
    {"type": "function", "function": {
        "name": "fetch_url",
        "description": "Открывает веб-страницу по ссылке и возвращает её текст. Используй, чтобы прочитать статью, документацию или результат поиска подробнее.",
        "parameters": {"type": "object", "properties": {"url": {"type": "string"}}, "required": ["url"]}}},
]


def tool_prompt():
    now = datetime.now(timezone(timedelta(hours=3))).strftime("%d.%m.%Y %H:%M")
    return ("Ты — ассистент внутри Telegram-приложения AI Studio. Сейчас " + now + " (МСК).\n"
            "У тебя есть инструменты:\n"
            "• create_files — реально создаёт файлы/архивы и отправляет их пользователю. Если просят файл, документ, "
            "таблицу, код в файле или архив — ОБЯЗАТЕЛЬНО вызови create_files. Никогда не говори, что не можешь создать "
            "или отправить файл.\n"
            "• web_search и fetch_url — доступ к интернету. Используй для актуальной информации и ссылок; указывай источники.\n"
            "После вызова инструмента коротко сообщи результат. Отвечай на языке пользователя.")


def tool_label(name, args):
    if name == "create_files":
        if args.get("archive_name"):
            return "Создаю архив " + str(args["archive_name"])
        n = len(args.get("files") or [])
        return "Создаю файл " + str((args.get("files") or [{}])[0].get("path", "")) if n == 1 else f"Создаю файлы ({n})"
    if name == "web_search":
        return "Ищу в интернете: " + str(args.get("query", ""))[:60]
    if name == "fetch_url":
        return "Читаю " + (urlparse(str(args.get("url", ""))).netloc or "страницу")
    return name


# ---------- files ----------
def safe_part(p):
    p = re.sub(r'[\x00-\x1f<>:"|?*\\]', "_", p).strip(" .")
    return p[:120] or "file"


def safe_path(path):
    parts = [safe_part(x) for x in str(path).replace("\\", "/").split("/") if x not in ("", ".", "..")]
    return "/".join(parts[-6:]) or "file.txt"


def md_to_docx(text):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    def add_runs(par, line):
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", line)):
            if part:
                par.add_run(part).bold = bool(i % 2)

    in_code = False
    for raw in (text or "").splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            r = doc.add_paragraph().add_run(line)
            r.font.name = "Consolas"
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
        elif re.match(r"^\s*[-*•]\s+", line):
            add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*•]\s+", "", line))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line))
        elif line.strip():
            add_runs(doc.add_paragraph(), line)
    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()


def csv_to_xlsx(text):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    sheets, cur, name = [], [], "Лист1"
    for line in (text or "").splitlines():
        m = re.match(r"^###\s*Sheet:\s*(.+)", line.strip(), re.I)
        if m:
            if cur:
                sheets.append((name, cur))
            name, cur = m.group(1).strip()[:31], []
        else:
            cur.append(line)
    sheets.append((name, cur))
    for name, lines in sheets:
        ws = wb.create_sheet(safe_part(name)[:31] or "Лист")
        body = "\n".join(lines).strip()
        if not body:
            continue
        try:
            dialect = csv.Sniffer().sniff(body[:2000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        for r, row in enumerate(csv.reader(io.StringIO(body), dialect), 1):
            for c, val in enumerate(row, 1):
                v = val.strip()
                try:
                    v = float(v.replace(",", ".")) if re.fullmatch(r"-?\d+([.,]\d+)?", v) else v
                    v = int(v) if isinstance(v, float) and v.is_integer() and "." not in val and "," not in val else v
                except ValueError:
                    pass
                cell = ws.cell(row=r, column=c, value=v)
                if r == 1:
                    cell.font = Font(bold=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, *(len(str(c.value or "")) + 2 for c in col)))
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def build_file(path, content):
    ext = path.lower().rsplit(".", 1)[-1] if "." in path else ""
    if ext == "docx":
        return md_to_docx(content)
    if ext == "xlsx":
        return csv_to_xlsx(content)
    return (content or "").encode("utf-8")


def cleanup_files():
    if not os.path.isdir(FILES_DIR):
        return
    for d in os.listdir(FILES_DIR):
        p = os.path.join(FILES_DIR, d)
        try:
            if time.time() - os.path.getmtime(p) > FILE_TTL:
                for f in os.listdir(p):
                    os.remove(os.path.join(p, f))
                os.rmdir(p)
        except OSError:
            pass


def save_output(name, data, public_url):
    fid = uuid.uuid4().hex
    d = os.path.join(FILES_DIR, fid)
    os.makedirs(d, exist_ok=True)
    with open(os.path.join(d, name), "wb") as f:
        f.write(data)
    return {"name": name, "size": len(data), "url": f"{public_url}/files/{fid}/{quote(name)}", "_data": data}


def create_files(args, public_url):
    """Возвращает (текст для модели, список файлов)."""
    cleanup_files()
    files = args.get("files") or []
    if isinstance(files, dict):
        files = [files]
    if not files:
        return "Ошибка: не передано ни одного файла.", []
    built, total = [], 0
    for f in files[:100]:
        path = safe_path(f.get("path") or "file.txt")
        data = build_file(path, f.get("content") if isinstance(f.get("content"), str) else json.dumps(f.get("content"), ensure_ascii=False))
        total += len(data)
        if total > MAX_TOTAL:
            return "Ошибка: слишком большой объём файлов (максимум 15 МБ).", []
        built.append((path, data))
    out = []
    arch = args.get("archive_name")
    if arch:
        arch = safe_part(str(arch).split("/")[-1])
        if not arch.lower().endswith(".zip"):
            arch += ".zip"
        b = io.BytesIO()
        with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
            for path, data in built:
                z.writestr(path, data)
        out.append(save_output(arch, b.getvalue(), public_url))
    else:
        for path, data in built:
            out.append(save_output(path.split("/")[-1], data, public_url))
    names = ", ".join(f"{o['name']} ({o['size']} байт)" for o in out)
    inside = f" Внутри архива: {', '.join(p for p, _ in built)}." if arch else ""
    return f"Готово. Файлы созданы и уже отправлены пользователю: {names}.{inside}", out


# ---------- web ----------
def strip_html(h):
    h = re.sub(r"(?is)<(script|style|noscript|svg|head|nav|footer|form)[^>]*>.*?</\1>", " ", h)
    h = re.sub(r"(?i)<br\s*/?>|</(p|div|li|h[1-6]|tr|section|article)>", "\n", h)
    h = re.sub(r"<[^>]+>", " ", h)
    h = htmlmod.unescape(h)
    h = re.sub(r"[ \t\r\f\v]+", " ", h)
    return re.sub(r"\n\s*\n+", "\n\n", h).strip()


def parse_ddg(page):
    res = []
    for block in re.split(r'<div[^>]+class="[^"]*result results_links', page)[1:]:
        a = re.search(r'<a[^>]*class="result__a"[^>]*href="([^"]+)"[^>]*>(.*?)</a>', block, re.S) or \
            re.search(r'<a[^>]*href="([^"]+)"[^>]*class="result__a"[^>]*>(.*?)</a>', block, re.S)
        if not a:
            continue
        href = htmlmod.unescape(a.group(1))
        if "uddg=" in href:
            href = parse_qs(urlparse(href if href.startswith("http") else "https:" + href).query).get("uddg", [href])[0]
        if "duckduckgo.com/y.js" in href:  # реклама
            continue
        sn = re.search(r'class="result__snippet"[^>]*>(.*?)</(a|div)>', block, re.S)
        res.append({"title": strip_html(a.group(2)), "url": href, "snippet": strip_html(sn.group(1)) if sn else ""})
    return res


def parse_lite(page):
    res = []
    for m in re.finditer(r'<a[^>]+href="([^"]+)"[^>]*class=[\'"]result-link[\'"][^>]*>(.*?)</a>(.*?)(?=class=[\'"]result-link|$)', page, re.S):
        href = htmlmod.unescape(m.group(1))
        if "uddg=" in href:
            href = parse_qs(urlparse(href if href.startswith("http") else "https:" + href).query).get("uddg", [href])[0]
        sn = re.search(r"class=['\"]result-snippet['\"][^>]*>(.*?)</td>", m.group(3), re.S)
        res.append({"title": strip_html(m.group(2)), "url": href, "snippet": strip_html(sn.group(1)) if sn else ""})
    return res


async def web_search(http, query):
    import aiohttp
    query = str(query)[:300]
    attempts = [("https://html.duckduckgo.com/html/", parse_ddg), ("https://lite.duckduckgo.com/lite/", parse_lite)]
    for url, parser in attempts:
        try:
            async with http.post(url, data={"q": query, "kl": "ru-ru"}, headers={"User-Agent": UA, "Accept-Language": "ru,en;q=0.8"},
                                 timeout=aiohttp.ClientTimeout(total=15)) as r:
                page = await r.text(errors="ignore")
            items = parser(page)[:8]
            if items:
                return "\n\n".join(f"{i}. {x['title']}\n{x['url']}\n{x['snippet']}" for i, x in enumerate(items, 1))
        except Exception:
            continue
    return "Поиск сейчас недоступен или ничего не нашёл. Ответь по своим знаниям и предупреди, что данные могут быть неактуальны."


async def is_public_host(host):
    try:
        infos = await asyncio.get_running_loop().getaddrinfo(host, None)
        return all(ipaddress.ip_address(i[4][0]).is_global for i in infos)
    except Exception:
        return False


async def fetch_url(http, url):
    import aiohttp
    url = str(url).strip()
    if not url.startswith(("http://", "https://")):
        url = "https://" + url
    host = urlparse(url).hostname or ""
    if not host or not await is_public_host(host):
        return "Ошибка: адрес недоступен."
    try:
        async with http.get(url, headers={"User-Agent": UA, "Accept-Language": "ru,en;q=0.8"}, allow_redirects=True,
                            timeout=aiohttp.ClientTimeout(total=20), max_redirects=5) as r:
            raw = await r.content.read(3 * 1024 * 1024)
            ctype = r.headers.get("Content-Type", "")
            status = r.status
        text = raw.decode(r.charset or "utf-8", errors="ignore")
        if "html" in ctype or text.lstrip()[:15].lower().startswith(("<!doctype", "<html")):
            title = re.search(r"(?is)<title[^>]*>(.*?)</title>", text)
            text = (strip_html(title.group(1)) + "\n\n" if title else "") + strip_html(text)
        return f"[HTTP {status}] {url}\n\n{text[:14000]}"
    except Exception as e:
        return f"Не удалось открыть страницу: {e}"


async def run_tool(http, name, args, public_url):
    """Возвращает (результат для модели, файлы)."""
    try:
        if name == "create_files":
            return create_files(args, public_url)
        if name == "web_search":
            return await web_search(http, args.get("query", "")), []
        if name == "fetch_url":
            return await fetch_url(http, args.get("url", "")), []
        return f"Неизвестный инструмент {name}", []
    except Exception as e:
        return f"Ошибка инструмента {name}: {e}", []
