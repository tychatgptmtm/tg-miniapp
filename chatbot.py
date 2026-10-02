"""Чат с нейросетью прямо в Telegram — без мини-приложения."""
import asyncio, base64, io, json, logging, re, time, zipfile

from tgformat import md_to_html, split_md, strip_md, esc

log = logging.getLogger("chat")
D = {}            # зависимости из server.py: tg, tg_upload, download, agent, store, models, admins, default_model, public_url
HIST: dict = {}   # uid -> история диалога (последние сообщения)
BUSY: set = set()
MAX_HIST = 24
TEXT_EXT = {"txt", "md", "py", "js", "ts", "jsx", "tsx", "json", "csv", "tsv", "html", "htm", "css", "xml", "yaml", "yml",
            "ini", "cfg", "toml", "log", "sql", "sh", "bat", "ps1", "java", "kt", "c", "h", "cpp", "hpp", "cs", "go", "rs",
            "php", "rb", "swift", "lua", "r", "env", "srt", "tex"}


def init(**deps):
    D.update(deps)


def tg(method, **p):
    return D["tg"](method, **p)


# ---------- модель пользователя ----------
def user_model(uid):
    return (D["store"].users.get(str(uid)) or {}).get("model") or D["default_model"]


def set_user_model(uid, model):
    u = D["store"].users.get(str(uid))
    if u is not None:
        u["model"] = model
        D["store"].mark(soon=True)


def short_name(name):
    return re.sub(r"^[^:]+:\s*", "", name).replace(" (free)", "")


# ---------- входящие сообщения ----------
async def file_bytes(file_id, max_size=20 * 1024 * 1024):
    f = await tg("getFile", file_id=file_id)
    if not f.get("ok"):
        raise ValueError("Не удалось получить файл (Telegram отдаёт ботам файлы до 20 МБ)")
    if (f["result"].get("file_size") or 0) > max_size:
        raise ValueError("Файл слишком большой")
    return await D["download"](f["result"]["file_path"])


def doc_text(name, data):
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext in TEXT_EXT or not ext:
        for enc in ("utf-8", "cp1251"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                pass
        return data.decode("utf-8", "replace")
    if ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            parts += [" | ".join(c.text for c in r.cells) for r in t.rows]
        return "\n".join(parts)
    if ext in ("xlsx", "xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"### Лист: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    out.append(",".join("" if v is None else str(v) for v in row))
                if len(out) > 3000:
                    break
        return "\n".join(out)
    if ext == "pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            return None
        r = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in r.pages[:60])
    if ext == "zip":
        z, out = zipfile.ZipFile(io.BytesIO(data)), []
        for info in z.infolist()[:200]:
            out.append(f"--- {info.filename} ({info.file_size} байт)")
            sub = info.filename.rsplit(".", 1)[-1].lower()
            if not info.is_dir() and sub in TEXT_EXT and info.file_size < 100_000:
                out.append(z.read(info).decode("utf-8", "replace"))
        return "\n".join(out)
    return None


async def build_content(msg):
    """Возвращает (content для модели, текст для истории) или (None, причина)."""
    text = msg.get("text") or msg.get("caption") or ""
    img = None
    if msg.get("photo"):
        img = (await file_bytes(msg["photo"][-1]["file_id"]), "image/jpeg")
    elif doc := msg.get("document"):
        name, mime = doc.get("file_name") or "file", doc.get("mime_type") or ""
        data = await file_bytes(doc["file_id"])
        if mime.startswith("image/"):
            img = (data, mime)
        else:
            body = doc_text(name, data)
            if body is None:
                return None, f"Не умею читать файлы такого типа ({esc(name)}). Пришлите текст, фото, PDF, Word, Excel, ZIP или файл с кодом."
            body = body[:60000] + ("\n…(обрезано)" if len(body) > 60000 else "")
            full = f"{text}\n\nФайл «{name}»:\n```\n{body}\n```".strip()
            return full, f"{text}\n[пользователь прислал файл «{name}»]".strip()
    elif msg.get("voice") or msg.get("audio") or msg.get("video_note"):
        return None, "Голосовые пока не понимаю — напишите, пожалуйста, текстом."
    elif msg.get("sticker"):
        return None, "Классный стикер! Но я понимаю текст, фото и файлы 🙂"
    elif msg.get("video") or msg.get("animation"):
        return None, "Видео пока не понимаю — пришлите кадр фотографией или опишите текстом."
    if img:
        url = f"data:{img[1]};base64," + base64.b64encode(img[0]).decode()
        content = [{"type": "text", "text": text or "Что на этом изображении? Опиши подробно."},
                   {"type": "image_url", "image_url": {"url": url}}]
        return content, f"{text}\n[пользователь прислал изображение]".strip()
    if not text.strip():
        return None, "Напишите сообщение, и я отвечу."
    return text, text


# ---------- вывод ответа с «живым» обновлением ----------
class Reply:
    def __init__(self, chat_id, reply_to=None):
        self.chat_id, self.reply_to = chat_id, reply_to
        self.ids, self.sent = [], []
        self.text, self.status, self.last = "", "Думаю…", 0.0

    async def _put(self, i, html, plain):
        if i < len(self.ids):
            if self.sent[i] == html:
                return
            r = await tg("editMessageText", chat_id=self.chat_id, message_id=self.ids[i], text=html, parse_mode="HTML",
                         link_preview_options={"is_disabled": True})
            if not r.get("ok") and "not modified" not in str(r.get("description")):
                await tg("editMessageText", chat_id=self.chat_id, message_id=self.ids[i], text=plain[:4096],
                         link_preview_options={"is_disabled": True})
            self.sent[i] = html
        else:
            extra = {"reply_parameters": {"message_id": self.reply_to, "allow_sending_without_reply": True}} if self.reply_to and i == 0 else {}
            r = await tg("sendMessage", chat_id=self.chat_id, text=html, parse_mode="HTML", link_preview_options={"is_disabled": True}, **extra)
            if not r.get("ok"):
                r = await tg("sendMessage", chat_id=self.chat_id, text=plain[:4096] or "…")
            self.ids.append((r.get("result") or {}).get("message_id"))
            self.sent.append(html)

    async def render(self, force=False):
        if not force and time.monotonic() - self.last < 1.4:
            return
        self.last = time.monotonic()
        parts = split_md(self.text.strip()) if self.text.strip() else [""]
        for i, part in enumerate(parts):
            html, plain = md_to_html(part), strip_md(part)
            if i == len(parts) - 1 and self.status:
                html = (html + "\n\n" if html else "") + f"<i>{esc(self.status)}</i>"
                plain = (plain + "\n\n" if plain else "") + self.status
            await self._put(i, html or "…", plain or "…")


async def keep_typing(chat_id, stop: asyncio.Event):
    while not stop.is_set():
        await tg("sendChatAction", chat_id=chat_id, action="typing")
        try:
            await asyncio.wait_for(stop.wait(), 4.5)
        except asyncio.TimeoutError:
            pass


def friendly(m):
    if re.search(r"data policy|No endpoints found matching", m, re.I):
        return m + "\n\nВключите бесплатные модели: openrouter.ai → Settings → Privacy."
    if re.search(r"rate.?limit|429", m, re.I):
        return "Слишком много запросов к модели (лимит бесплатного тарифа). Подождите минуту или смените модель: /model"
    if re.search(r"credits|402|insufficient", m, re.I):
        return "Недостаточно средств на OpenRouter для этой модели. Выберите бесплатную: /model"
    if re.search(r"image|vision|multimodal", m, re.I):
        return "Эта модель не понимает изображения. Выберите другую: /model"
    return m


async def answer(chat_id, user, msg):
    uid = user["id"]
    if uid in BUSY:
        await tg("sendMessage", chat_id=chat_id, text="Секунду, ещё отвечаю на прошлое сообщение…")
        return
    BUSY.add(uid)
    stop = asyncio.Event()
    typing = asyncio.get_running_loop().create_task(keep_typing(chat_id, stop))
    try:
        try:
            content, hist_text = await build_content(msg)
        except Exception as e:
            content, hist_text = None, f"Не получилось прочитать: {esc(str(e))}"
        if content is None:
            await tg("sendMessage", chat_id=chat_id, text=hist_text, parse_mode="HTML")
            return
        D["store"].touch(user, msg=True)
        hist = HIST.setdefault(uid, [])
        body = {"model": user_model(uid), "messages": hist[-MAX_HIST:] + [{"role": "user", "content": content}],
                "temperature": 0.7, "tools": True}
        rep = Reply(chat_id, msg.get("message_id"))
        await rep.render(force=True)
        files, err = [], None
        gen = D["agent"](body, user)
        try:
            async for ev in gen:
                if ev.get("error"):
                    err = (ev["error"] or {}).get("message") or "Ошибка модели"
                    break
                if t := ev.get("x_tool"):
                    rep.status = (t.get("label") or "Работаю") + "…" if t.get("status") == "start" else "Думаю…"
                    await rep.render(force=True)
                    continue
                if f := ev.get("x_file"):
                    files.append(f["name"])
                    continue
                ch = (ev.get("choices") or [{}])[0]
                delta = ch.get("delta") or ch.get("message") or {}
                c = delta.get("content")
                if isinstance(c, list):
                    c = "".join(p.get("text", "") for p in c if p.get("type") == "text")
                if c:
                    rep.text += c
                    if rep.text.strip():
                        rep.status = ""
                    await rep.render()
        finally:
            await gen.aclose()
        rep.status = ""
        if err:
            rep.text = (rep.text.rstrip() + "\n\n" if rep.text.strip() else "") + "⚠️ " + friendly(err)
        elif not rep.text.strip():
            rep.text = "Готово — файлы выше." if files else "Пустой ответ. Попробуйте ещё раз или смените модель: /model"
        await rep.render(force=True)
        if not err:
            note = f"\n\n[Созданы и отправлены пользователю файлы: {', '.join(files)}]" if files else ""
            hist += [{"role": "user", "content": hist_text}, {"role": "assistant", "content": (rep.text + note).strip()}]
            del hist[:-MAX_HIST]
    except Exception as e:
        log.exception("answer failed")
        await tg("sendMessage", chat_id=chat_id, text=f"Что-то пошло не так: {e}")
    finally:
        stop.set()
        typing.cancel()
        BUSY.discard(uid)


# ---------- выбор модели ----------
MENU: list = []


async def model_menu(chat_id, uid, edit_id=None):
    models = await D["models"]()
    cur, default = user_model(uid), D["default_model"]
    free = [m for m in models if m.get("free") and m.get("textOut", True) and m["id"] != default]
    free.sort(key=lambda m: -(m.get("created") or 0))
    MENU.clear()
    MENU.extend([{"id": default, "name": next((m["name"] for m in models if m["id"] == default), default)}] + free[:13])
    rows = [[{"text": ("● " if m["id"] == cur else "") + short_name(m["name"])[:40], "callback_data": f"md:{i}"}] for i, m in enumerate(MENU)]
    text = (f"<b>Модель сейчас:</b> <code>{esc(cur)}</code>\n\nВыберите бесплатную модель кнопкой ниже "
            "или пришлите команду с ID модели с openrouter.ai/models, например:\n<code>/model openai/gpt-4o-mini</code>")
    if edit_id:
        await tg("editMessageText", chat_id=chat_id, message_id=edit_id, text=text, parse_mode="HTML", reply_markup={"inline_keyboard": rows})
    else:
        await tg("sendMessage", chat_id=chat_id, text=text, parse_mode="HTML", reply_markup={"inline_keyboard": rows})


# ---------- админка в чате ----------
def ago(t):
    s = time.time() - (t or 0)
    return "сейчас" if s < 120 else f"{int(s // 60)} мин" if s < 3600 else f"{int(s // 3600)} ч" if s < 86400 else f"{int(s // 86400)} дн"


async def admin_panel(chat_id, edit_id=None, page=0):
    st, store = D["store"].stats(), D["store"]
    users = sorted(store.users.values(), key=lambda u: u.get("last_seen", 0), reverse=True)
    per = 8
    chunk = users[page * per:(page + 1) * per]
    lines = [f"<b>Пользователи</b>\n\nВсего: <b>{st['total']}</b> · за 24 ч: <b>{st['active24']}</b>\n"
             f"Сообщений сегодня: <b>{st['msgsToday']}</b> · заблокировано: <b>{st['blocked']}</b>\n"]
    rows = []
    for u in chunk:
        nm = u.get("name") or (("@" + u["username"]) if u.get("username") else str(u["id"]))
        tag = " · админ" if u["id"] in D["admins"] else " · ЗАБЛОКИРОВАН" if u.get("blocked") else ""
        lines.append(f"{'⛔' if u.get('blocked') else '•'} <b>{esc(nm)}</b>{(' @' + esc(u['username'])) if u.get('username') else ''}"
                     f" — {ago(u.get('last_seen'))} назад, {u.get('msgs', 0)} сообщ.{tag}")
        if u["id"] not in D["admins"]:
            rows.append([{"text": ("Разблокировать · " if u.get("blocked") else "Заблокировать · ") + nm[:28],
                          "callback_data": f"ub:{u['id']}:{0 if u.get('blocked') else 1}:{page}"}])
    nav = []
    if page > 0:
        nav.append({"text": "← Назад", "callback_data": f"ap:{page - 1}"})
    nav.append({"text": "Обновить", "callback_data": f"ap:{page}"})
    if (page + 1) * per < len(users):
        nav.append({"text": "Дальше →", "callback_data": f"ap:{page + 1}"})
    rows.append(nav)
    text, kb = "\n".join(lines), {"inline_keyboard": rows}
    if edit_id:
        await tg("editMessageText", chat_id=chat_id, message_id=edit_id, text=text, parse_mode="HTML", reply_markup=kb)
    else:
        await tg("sendMessage", chat_id=chat_id, text=text, parse_mode="HTML", reply_markup=kb)


async def on_callback(cq):
    """True, если callback обработан здесь."""
    data, uid = cq.get("data") or "", cq["from"]["id"]
    chat_id, mid = cq["message"]["chat"]["id"], cq["message"]["message_id"]
    if data.startswith("md:"):
        i = int(data[3:])
        if i < len(MENU):
            set_user_model(uid, MENU[i]["id"])
            await tg("answerCallbackQuery", callback_query_id=cq["id"], text="Модель: " + short_name(MENU[i]["name"]))
            await model_menu(chat_id, uid, edit_id=mid)
        else:
            await tg("answerCallbackQuery", callback_query_id=cq["id"], text="Меню устарело — отправьте /model ещё раз")
        return True
    if data.startswith(("ub:", "ap:")):
        if uid not in D["admins"]:
            await tg("answerCallbackQuery", callback_query_id=cq["id"], text="Только для администратора")
            return True
        if data.startswith("ub:"):
            _, target, block, page = data.split(":")
            target, block = int(target), block == "1"
            D["store"].set_blocked(target, block)
            await tg("sendMessage", chat_id=target, text="Администратор ограничил вам доступ к боту." if block else "Доступ к боту снова открыт.")
            await tg("answerCallbackQuery", callback_query_id=cq["id"], text="Заблокирован" if block else "Разблокирован")
        else:
            page = data[3:]
            await tg("answerCallbackQuery", callback_query_id=cq["id"])
        await admin_panel(chat_id, edit_id=mid, page=int(page))
        return True
    return False


async def on_command(cmd, arg, chat_id, user):
    """True, если команда обработана здесь."""
    uid = user["id"]
    if cmd == "/new":
        HIST.pop(uid, None)
        await tg("sendMessage", chat_id=chat_id, text="Начали новый диалог. О чём поговорим?")
        return True
    if cmd == "/model":
        if arg:
            set_user_model(uid, arg.strip())
            await tg("sendMessage", chat_id=chat_id, parse_mode="HTML", text=f"Готово, модель: <code>{esc(arg.strip())}</code>")
        else:
            await model_menu(chat_id, uid)
        return True
    if cmd in ("/admin", "/users"):
        if uid in D["admins"]:
            await admin_panel(chat_id)
        else:
            await tg("sendMessage", chat_id=chat_id, text="Эта команда только для администратора.")
        return True
    return False
